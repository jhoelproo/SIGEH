"""Export confirmation persists atomically on a disposable PostgreSQL server."""

from concurrent.futures import ThreadPoolExecutor
import hashlib
import json

from openpyxl import load_workbook
import pytest

import CALCULOS_QT as app
from private_insurance_exporter import create_private_ars_workbook
from tests.test_billing_consistency_postgres import server as server
from tests.test_receipt_optional_uuid_postgres import receipts as receipts, save
from tests.test_receipt_list_consistency_postgres import USER, batch, correct


@pytest.fixture
def exported(receipts, tmp_path):
    receipt_id = save()
    batch_id = batch(receipt_id)
    correct(batch_id, receipt_id)
    app.update_monthly_batch_configuration(
        batch_id,
        {
            "ars_display_name": "FUTURO",
            "invoice_number": "FAC-PRUEBA-001",
            "ncf": "B0100000001",
            "invoice_date": "2026-09-30",
            "provider_code": "14427",
            "provider_name": "HOSPITAL SINTETICO",
            "provider_rnc": "430130516",
            "ars_rnc": "101000001",
            "ars_address": "DIRECCION DE PRUEBA",
            "director_name": "DIRECTORA DE PRUEBA",
        },
        USER,
    )
    with app.db_connect() as con:
        con.execute(
            "UPDATE recibos SET estado_facturacion='FACTURADO' WHERE id=%s",
            (receipt_id,),
        )
    path = tmp_path / "listado_final.xlsx"
    create_private_ars_workbook(
        path,
        app.get_monthly_billing_batch(batch_id),
        app.list_monthly_batch_receipts(batch_id),
    )
    app.record_monthly_batch_export(batch_id, str(path), USER)
    return receipt_id, batch_id, path


def sent_events(batch_id):
    with app.db_connect() as con:
        rows = con.execute(
            """SELECT details FROM billing_batch_events
               WHERE batch_id=%s AND event_type='CONFIRMADO_ENVIADO'""",
            (batch_id,),
        ).fetchall()
    return [json.loads(row["details"]) for row in rows]


@pytest.mark.parametrize("edited", [False, True])
def test_final_export_is_sent_once_with_frozen_receipts_and_file_audit(
    exported, edited
):
    receipt_id, batch_id, path = exported
    original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    before = app.get_recibo_data(receipt_id)
    if edited:
        workbook = load_workbook(path)
        workbook["Relación Emergencias"]["E15"] = "00009999"
        workbook.save(path)
        workbook.close()

    result = app.confirm_monthly_batch_sent(batch_id, USER)

    assert result["receipts"] == 1
    assert result["total"] == 100
    reopened = app.get_monthly_billing_batch(batch_id)
    assert reopened["status"] == app.BATCH_SENT
    assert reopened["sent_receipt_count"] == 1
    assert float(reopened["sent_total"]) == 100
    event = sent_events(batch_id)
    assert len(event) == 1
    assert event[0]["excel_modified"] == edited
    assert event[0]["exported_file_sha256"] == original_hash
    assert event[0]["sent_file_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert app.get_recibo_data(receipt_id) == before
    assert len(app.list_monthly_batch_receipts(batch_id)) == 1
    with pytest.raises(ValueError, match="PENDIENTE"):
        app.confirm_monthly_batch_sent(batch_id, USER)
    assert len(sent_events(batch_id)) == 1
    with pytest.raises(ValueError):
        app.update_monthly_batch_configuration(batch_id, {"ncf": "OTRO"}, USER)
    with app.db_connect() as con:
        assert con.execute("SELECT COUNT(*) FROM recibos").fetchone()[0] == 1
        assert (
            con.execute("SELECT COUNT(*) FROM billing_batch_receipts").fetchone()[0]
            == 1
        )
        con.execute("UPDATE recibos SET total=9999 WHERE id=%s", (receipt_id,))
    assert (
        app.list_sent_billing_batches_report(batch_ids=[batch_id])[0]["sent_total"]
        == 100
    )


def test_batch_changes_are_not_mistaken_for_external_excel_edits(exported):
    _receipt_id, batch_id, _path = exported
    with app.db_connect() as con:
        con.execute(
            "UPDATE billing_batches SET invoice_number='FAC-NUEVA' WHERE id=%s",
            (batch_id,),
        )
    with pytest.raises(ValueError, match="cambió después"):
        app.confirm_monthly_batch_sent(batch_id, USER)
    assert app.get_monthly_billing_batch(batch_id)["status"] == app.BATCH_PENDING
    assert sent_events(batch_id) == []


def test_audit_failure_rolls_back_send_and_retry_creates_one_event(
    exported, monkeypatch
):
    _receipt_id, batch_id, _path = exported
    history = app._insert_action_history

    def reject(*_args, **_kwargs):
        raise RuntimeError("Fallo simulado de auditoría")

    monkeypatch.setattr(app, "_insert_action_history", reject)
    with pytest.raises(RuntimeError, match="auditoría"):
        app.confirm_monthly_batch_sent(batch_id, USER)
    assert app.get_monthly_billing_batch(batch_id)["status"] == app.BATCH_PENDING
    assert sent_events(batch_id) == []
    monkeypatch.setattr(app, "_insert_action_history", history)
    app.confirm_monthly_batch_sent(batch_id, USER)
    assert len(sent_events(batch_id)) == 1


def test_two_stations_cannot_confirm_the_same_batch_twice(exported):
    _receipt_id, batch_id, path = exported
    workbook = load_workbook(path)
    workbook["Relación Emergencias"]["E15"] = "00009999"
    workbook.save(path)
    workbook.close()

    def confirm():
        try:
            app.confirm_monthly_batch_sent(batch_id, USER)
            return "SENT"
        except ValueError as exc:
            assert "PENDIENTE" in str(exc)
            return "ALREADY_SENT"

    with ThreadPoolExecutor(max_workers=2) as workers:
        results = list(workers.map(lambda _: confirm(), range(2)))
    assert sorted(results) == ["ALREADY_SENT", "SENT"]
    assert len(sent_events(batch_id)) == 1
