"""Word imports follow the document's catalog rather than the active tab."""

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document

import CALCULOS_QT as app


PRICES = Path(__file__).parent / "fixtures/word/catalogo_sintetico.docx"


@pytest.fixture
def imports(monkeypatch):
    universal = {cat: {"inserted": 1, "updated": 0} for cat in app.UNIVERSAL_CATEGORIES}
    ars = {cat: {"inserted": 1, "updated": 0} for cat in app.ARS_CATEGORIES}
    spies = SimpleNamespace(
        universal=Mock(return_value=universal),
        ars=Mock(return_value=ars),
        information=Mock(),
        critical=Mock(),
    )
    monkeypatch.setattr(app, "import_word_to_universal_catalog", spies.universal)
    monkeypatch.setattr(app, "import_word_to_ars_catalog", spies.ars)
    monkeypatch.setattr(
        app.QFileDialog, "getOpenFileName", lambda *_: (str(PRICES), "")
    )
    monkeypatch.setattr(app.QMessageBox, "information", spies.information)
    monkeypatch.setattr(app.QMessageBox, "critical", spies.critical)
    monkeypatch.setattr(app, "get_universal", lambda cat: {cat: 10})
    monkeypatch.setattr(app, "get_ars_items", lambda cat, ars: {cat: 20})
    monkeypatch.setattr(app, "log_action", Mock())
    monkeypatch.setattr(app.ARS_RUNTIME_CACHE, "invalidate", Mock())
    return spies


@pytest.mark.parametrize("category", app.ALL_CATEGORIES)
def test_medications_document_imports_from_every_billing_tab(category, imports):
    window = SimpleNamespace(
        mark_activity=Mock(),
        get_current_category=lambda: category,
        current_ars="ARS QA",
        current_user={"username": "QA"},
        universal={},
        ars_cache={},
        refresh_picker=Mock(),
        search_and_maybe_switch_tab=Mock(),
    )
    app.MainWindow.import_meds_mats_from_word(window)
    imports.universal.assert_called_once_with(str(PRICES))
    imports.ars.assert_not_called()
    imports.critical.assert_not_called()
    assert "Medicamentos" in imports.information.call_args.args[2]
    assert window.universal == {
        "Medicamentos": {"Medicamentos": 10},
        "Materiales": {"Materiales": 10},
    }
    window.refresh_picker.assert_called_once()


def test_medications_document_imports_from_ars_manager(imports):
    window = SimpleNamespace(
        ars_combo=SimpleNamespace(currentText=lambda: "ARS QA"),
        current_user={"username": "QA"},
    )
    app.ARSManagerDialog.import_ars_word(window)
    imports.universal.assert_called_once_with(str(PRICES))
    imports.ars.assert_not_called()
    imports.critical.assert_not_called()
    assert "Medicamentos" in imports.information.call_args.args[2]


def test_synthetic_prices_document_keeps_names_and_base_prices():
    parsed = app.parse_word_for_universal_categories(str(PRICES))
    assert len(parsed["Medicamentos"]) == 2
    assert len(parsed["Materiales"]) == 2
    meds, materials = (dict(parsed[cat]) for cat in app.UNIVERSAL_CATEGORIES)
    assert meds["PRODUCTO QA ALFA 10 MG"] == 12.34
    assert meds["PRODUCTO QA BETA 5 MG"] == 0.43
    assert materials["INSUMO QA JERINGA"] == 2.00
    assert materials["INSUMO QA GUANTES"] == 9180.00


def write_document(tmp_path, *paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    path = tmp_path / "catalog.docx"
    document.save(path)
    return str(path)


@pytest.mark.parametrize("category", app.ARS_CATEGORIES)
def test_service_document_routes_to_ars_from_medications_tab(
    tmp_path, imports, category
):
    path = write_document(tmp_path, category, "SERVICIO QA $100.00")
    target, _summary = app.import_word_catalog(path, "Medicamentos", "ARS QA")
    assert target == "ars"
    imports.ars.assert_called_once_with(path, "ARS QA")
    imports.universal.assert_not_called()


@pytest.mark.parametrize(
    "category,target", [("Medicamentos", "universal"), ("Honorarios", "ars")]
)
def test_no_heading_preserves_selected_catalog(tmp_path, imports, category, target):
    path = write_document(tmp_path, "SIN ENCABEZADO $12.00")
    assert app.import_word_catalog(path, category, "ARS QA")[0] == target


def test_medication_document_does_not_require_selected_ars(imports):
    assert app.import_word_catalog(str(PRICES), "Laboratorios", "")[0] == "universal"


def test_empty_ars_document_does_not_open_database(tmp_path, monkeypatch):
    path = write_document(tmp_path, "DOCUMENTO SIN PRECIOS")
    connect = Mock(side_effect=AssertionError("Unexpected write"))
    monkeypatch.setattr(app, "db_connect", connect)
    monkeypatch.setattr(
        app, "get_ars_items", Mock(side_effect=AssertionError("Unexpected read"))
    )
    with pytest.raises(ValueError, match="encabezados"):
        app.import_word_to_ars_catalog(path, "ARS QA")
    connect.assert_not_called()


def test_empty_universal_document_does_not_report_success(tmp_path, monkeypatch):
    path = write_document(tmp_path, "MEDICAMENTOS", "NINGUN PRECIO")
    connect = Mock(side_effect=AssertionError("Unexpected write"))
    monkeypatch.setattr(app, "db_connect", connect)
    monkeypatch.setattr(
        app, "get_universal", Mock(side_effect=AssertionError("Unexpected read"))
    )
    with pytest.raises(ValueError, match="precios"):
        app.import_word_to_universal_catalog(path)
    connect.assert_not_called()


def test_service_document_requires_ars_before_persistence(tmp_path, imports):
    path = write_document(tmp_path, "LABORATORIOS", "HEMOGRAMA $12.00")
    with pytest.raises(ValueError, match="Selecciona una ARS"):
        app.import_word_catalog(path, "Medicamentos", "")
    imports.ars.assert_not_called()


def test_missing_docx_dependency_is_explicit(imports, monkeypatch):
    monkeypatch.setattr(app, "Document", None)
    with pytest.raises(RuntimeError, match="python-docx"):
        app.import_word_catalog(str(PRICES), "Medicamentos", "ARS QA")
    imports.universal.assert_not_called()


def test_table_category_heading_selects_catalog(tmp_path, imports):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "MATERIALES"
    table.cell(1, 0).text = "JERINGA"
    table.cell(1, 1).text = "10.00"
    path = tmp_path / "table.docx"
    document.save(path)
    assert app.import_word_catalog(str(path), "Honorarios", "ARS QA")[0] == "universal"


def test_ars_manager_keeps_service_import_and_summary(tmp_path, imports, monkeypatch):
    path = write_document(tmp_path, "HONORARIOS", "HONORARIO DIA $100")
    monkeypatch.setattr(app.QFileDialog, "getOpenFileName", lambda *_: (path, ""))
    window = SimpleNamespace(
        ars_combo=SimpleNamespace(currentText=lambda: "ARS QA"),
        current_user={"username": "QA"},
    )
    app.ARSManagerDialog.import_ars_word(window)
    imports.ars.assert_called_once_with(path, "ARS QA")
    imports.universal.assert_not_called()
    assert "Honorarios" in imports.information.call_args.args[2]
    assert not getattr(window, "universal_catalog_changed", False)


def test_billing_reloads_service_catalog_after_detecting_ars_document(
    tmp_path, imports, monkeypatch
):
    path = write_document(tmp_path, "LABORATORIOS", "HEMOGRAMA $100")
    monkeypatch.setattr(app.QFileDialog, "getOpenFileName", lambda *_: (path, ""))
    window = SimpleNamespace(
        mark_activity=Mock(),
        get_current_category=lambda: "Medicamentos",
        current_ars="ARS QA",
        current_user={"username": "QA"},
        universal={},
        ars_cache={},
        refresh_picker=Mock(),
        search_and_maybe_switch_tab=Mock(),
    )
    app.MainWindow.import_meds_mats_from_word(window)
    imports.ars.assert_called_once_with(path, "ARS QA")
    assert set(window.ars_cache) == set(app.ARS_CATEGORIES)
    assert "ARS QA" in imports.information.call_args.args[2]
    window.refresh_picker.assert_called_once()


@pytest.mark.parametrize("changed", [False, True])
def test_closing_ars_manager_reloads_universal_only_when_changed(
    imports, monkeypatch, changed
):
    dialog = SimpleNamespace(
        exec=Mock(), _ars_names=["ARS QA"], universal_catalog_changed=changed
    )
    monkeypatch.setattr(app, "ARSManagerDialog", lambda *_: dialog)
    window = SimpleNamespace(
        mark_activity=Mock(),
        current_user={"username": "QA"},
        current_ars="ARS QA",
        universal={"old": {}},
        ars_combo=Mock(),
        on_ars_changed=Mock(),
    )
    window.ars_combo.currentText.return_value = "ARS QA"
    app.MainWindow.open_ars_manager(window)
    assert ("Medicamentos" in window.universal) == changed
    window.on_ars_changed.assert_called_once_with("ARS QA")


def test_failed_import_keeps_catalog_and_does_not_report_success(imports):
    imports.universal.side_effect = ValueError("Synthetic invalid document")
    window = SimpleNamespace(
        mark_activity=Mock(),
        get_current_category=lambda: "Medicamentos",
        current_ars="ARS QA",
        current_user={"username": "QA"},
        universal={"old": {}},
        ars_cache={},
        refresh_picker=Mock(),
    )
    app.MainWindow.import_meds_mats_from_word(window)
    imports.information.assert_not_called()
    imports.critical.assert_called_once()
    assert window.universal == {"old": {}}
    window.refresh_picker.assert_not_called()
