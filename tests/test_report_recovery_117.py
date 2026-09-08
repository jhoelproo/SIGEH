from datetime import datetime

import pytest

import CALCULOS_QT as app
from admission_statistical_reports import (
    AdmissionReportFilters,
    build_admission_report_dataset,
)


@pytest.mark.parametrize("clock", ["08:30 AM", "08:30:15 PM", "12:00 AM", "12:00 PM"])
def test_closed_turn_retains_all_legacy_civil_timestamps(clock):
    rows = [
        dict(
            attention_id=i,
            service_date="2026-09-07",
            service_time=clock,
            turn_id=3954,
            operational_source_id="source",
            patient_name="SYNTHETIC",
        )
        for i in range(1, 138)
    ]
    dataset = build_admission_report_dataset(
        rows,
        AdmissionReportFilters(
            start_at=datetime(2026, 9, 7, 8),
            end_at=datetime(2026, 9, 8, 9),
            turn_id=3954,
            operational_source_id="source",
            period_label="Synthetic turn",
        ),
    )
    assert dataset.summary["total_patients"] == 137


def test_closure_links_central_uuid_before_station_local_identity():
    records = [
        dict(
            source_instance_id="new-origin",
            attention_id=999,
            global_attention_id="central-a",
        )
    ]
    receipts = [
        dict(
            id=1,
            source_id="old-origin",
            admission_atencion_id=42,
            admission_global_attention_id="central-a",
            numero_autorizacion="123456",
        )
    ]
    assert (
        app._link_shift_receipts(records, receipts)[("new-origin", 999)] == receipts[0]
    )


def test_local_id_collision_cannot_link_another_central_attention():
    records = [
        dict(
            source_instance_id="origin",
            attention_id=42,
            global_attention_id="central-a",
        )
    ]
    receipts = [
        dict(
            id=1,
            source_id="origin",
            admission_atencion_id=42,
            admission_global_attention_id="central-b",
        )
    ]
    assert app._link_shift_receipts(records, receipts)[("origin", 42)] is None


@pytest.mark.parametrize(
    "variation,expected",
    [
        ("match", True),
        ("cedula", True),
        ("duplicate", False),
        ("date", False),
        ("ars", False),
        ("empty", False),
        ("linked", False),
        ("different", False),
    ],
)
def test_legacy_receipt_matching_remains_unambiguous(variation, expected):
    row = dict(
        source_instance_id="origin",
        attention_id=42,
        service_date="2026-09-07",
        ars="HUMANO",
        nss_snapshot="123",
        cedula_snapshot="456",
    )
    receipt = dict(id=1, fecha="2026-09-07", ars="HUMANO", admission_nss_snapshot="123")
    if variation == "date":
        receipt["fecha"] = "2026-09-06"
    if variation == "ars":
        receipt["ars"] = "OTHER"
    if variation == "empty":
        row.update(nss_snapshot="", cedula_snapshot="")
    if variation == "linked":
        receipt["admission_atencion_id"] = 99
    if variation in ("cedula", "different"):
        receipt.update(
            admission_nss_snapshot="999",
            admission_cedula_snapshot="456" if variation == "cedula" else "999",
        )
    receipts = [receipt, dict(receipt, id=2)] if variation == "duplicate" else [receipt]
    assert bool(app._link_shift_receipts([row], receipts)[("origin", 42)]) is expected


def test_long_excel_path_is_validated_and_copied_without_changing_original(tmp_path):
    from openpyxl import Workbook, load_workbook
    from excel_delivery_path import prepare_excel_delivery
    from pathlib import Path

    source = tmp_path / ("directory-" * 10) / ("listing-" * 15 + ".xlsx")
    source.parent.mkdir()
    workbook = Workbook()
    workbook.active["B6"] = "SYNTHETIC"
    workbook.save(source)
    before = source.read_bytes()
    result = Path(prepare_excel_delivery(source, tmp_path / "short"))
    assert len(str(result)) < len(str(source))
    assert source.read_bytes() == result.read_bytes() == before
    reopened = load_workbook(result)
    assert reopened.active["B6"].value == "SYNTHETIC"
    reopened.close()


@pytest.mark.parametrize("exists", [False, True])
def test_invalid_excel_is_not_sent_to_office(tmp_path, exists):
    from excel_delivery_path import prepare_excel_delivery
    from zipfile import BadZipFile

    source = tmp_path / "broken.xlsx"
    if exists:
        source.write_bytes(b"corruption evidence")
    with pytest.raises((FileNotFoundError, BadZipFile)):
        prepare_excel_delivery(source, tmp_path / "short")
    if exists:
        assert source.read_bytes() == b"corruption evidence"


def test_uncertain_print_does_not_block_startup(monkeypatch, tmp_path):
    from types import SimpleNamespace
    from unittest.mock import Mock
    from tests.test_turn_excel_post_commit import _v15_module
    import turn_excel_delivery as delivery

    v15 = _v15_module()
    queue = str(tmp_path / "queue.sqlite3")
    monkeypatch.setattr(v15, "EXCEL_PRINT_QUEUE_PATH", queue)
    delivery.enqueue(queue, {"transition_id": "unknown"}, 1)
    delivery.deliver(queue, "unknown", lambda _: "missing.xlsx", lambda *_: False)
    modal = Mock(side_effect=AssertionError("Startup must remain interactive"))
    monkeypatch.setattr(v15.messagebox, "showwarning", modal)
    page = SimpleNamespace(set_status=Mock(), _ejecutar_en_segundo_plano=Mock())
    v15.App._resume_turn_excel_delivery(page)
    assert page.set_status.called
    assert delivery.jobs(queue, "SUBMITTING")


def test_uncertain_artifact_repair_never_resubmits(tmp_path):
    from unittest.mock import Mock
    import turn_excel_delivery as delivery

    queue = tmp_path / "queue.sqlite3"
    delivery.enqueue(queue, {"transition_id": "old"}, 2)
    printer = Mock(return_value=False)
    delivery.deliver(queue, "old", lambda _: "missing.xlsx", printer)
    generate = Mock(return_value="recovered.xlsx")

    def valid(path):
        return path == "recovered.xlsx"

    assert delivery.repair_unconfirmed_artifacts(queue, generate, valid) == 1
    assert delivery.repair_unconfirmed_artifacts(queue, generate, valid) == 0
    assert delivery.jobs(queue, "SUBMITTING")[0]["file_path"] == "recovered.xlsx"
    assert printer.call_count == generate.call_count == 1


def test_rebuild_preserves_corruption_evidence_and_last_valid_file(tmp_path):
    from openpyxl import Workbook, load_workbook
    import excel_artifact as artifacts

    target = tmp_path / "list.xlsx"
    target.write_bytes(b"broken evidence")
    backup = tmp_path / "list.last-valid.xlsx"
    backup.write_bytes(b"previous evidence")
    book = Workbook()
    book.active["B6"] = "REBUILT FROM CANONICAL DATA"
    artifacts.save_workbook(book, target, recover_corrupt=True)
    assert artifacts.xlsx_is_valid(target)
    assert next(tmp_path.glob("*.corrupt-*.xlsx")).read_bytes() == b"broken evidence"
    assert backup.read_bytes() == b"previous evidence"
    loaded = load_workbook(target)
    assert loaded.active["B6"].value == "REBUILT FROM CANONICAL DATA"
    loaded.close()


def test_history_restores_same_window_after_minimizing():
    from tests.test_turn_excel_post_commit import _v15_module
    from PySide6.QtWidgets import QApplication, QLineEdit

    v15 = _v15_module()
    qt = QApplication.instance() or QApplication([])
    window = v15.Toplevel()
    search = QLineEdit(window)
    search.setText("FILTER PRESERVED")
    window.showMinimized()
    qt.processEvents()
    assert window.isMinimized()
    window.deiconify()
    qt.processEvents()
    try:
        assert not window.isMinimized()
        assert search.text() == "FILTER PRESERVED"
    finally:
        window.close()


@pytest.mark.parametrize("failure", ["queue", "context"])
def test_bad_print_queue_cannot_prevent_startup(monkeypatch, failure):
    from types import SimpleNamespace
    from unittest.mock import Mock
    from tests.test_turn_excel_post_commit import _v15_module
    import turn_excel_delivery

    v15 = _v15_module()
    if failure == "queue":
        reader = Mock(side_effect=OSError("synthetic queue failure"))
    else:
        reader = Mock(side_effect=[[{"context": {}}], []])
    monkeypatch.setattr(turn_excel_delivery, "jobs", reader)
    page = SimpleNamespace(set_status=Mock(), _run_turn_post_commit_effects=Mock())
    v15.App._resume_turn_excel_delivery(page)
    page._run_turn_post_commit_effects.assert_not_called()


def test_recovery_rebuilds_exact_context_without_printing(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import Mock
    from tests.test_turn_excel_post_commit import _v15_module
    import turn_excel_delivery

    v15 = _v15_module()
    context = object()
    monkeypatch.setattr(v15, "restore_turn_delivery_context", lambda data: context)
    snapshot = object()
    capture = Mock(return_value=snapshot)
    monkeypatch.setattr(v15, "build_turn_closure_report_snapshot", capture)
    generate = Mock(return_value=SimpleNamespace(excel_path="recovered.xlsx"))
    monkeypatch.setattr(v15, "generate_turn_closure_report_files", generate)

    def repair(path, generate, valid):
        assert generate({}) == "recovered.xlsx"
        return 1

    monkeypatch.setattr(turn_excel_delivery, "repair_unconfirmed_artifacts", repair)
    database = object()
    assert v15.repair_unconfirmed_turn_listings(database) == 1
    capture.assert_called_once_with(database, context)
    generate.assert_called_once_with(snapshot, generate_pdf=False)


def test_incomplete_report_snapshot_is_rejected(monkeypatch):
    from types import SimpleNamespace
    from tests.test_turn_excel_post_commit import _v15_module

    v15 = _v15_module()
    context = v15.OutgoingTurnContext(
        "source",
        17,
        1,
        1,
        "rep",
        "REP",
        datetime(2026, 9, 7, 8),
        datetime(2026, 9, 8, 8),
        "8AM_8AM",
        datetime(2026, 9, 7).date(),
        transition_id="transition",
        new_turn_id=18,
    )
    monkeypatch.setattr(
        v15, "_dataset_turno_central", lambda *a, **k: ([{}], 17, "source")
    )
    monkeypatch.setattr(
        v15,
        "build_admission_report_dataset",
        lambda *a, **k: SimpleNamespace(records=[]),
    )
    with pytest.raises(RuntimeError, match="no pudo interpretar"):
        v15.build_turn_closure_report_snapshot(object(), context)
