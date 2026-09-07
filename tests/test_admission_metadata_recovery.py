import json
import sqlite3
from types import SimpleNamespace

import pytest

from admission_v15_adapter import _HybridDatabaseProxy
from admission_specialty import resolve_specialty


@pytest.mark.parametrize("specialty", ["PEDIATRIA", "GINECOLOGIA", "GENERAL"])
def test_recovers_specialty_from_durable_event(specialty):
    for payload in (
        {"detail_sheet": specialty},
        json.dumps({"detail_sheet": specialty}),
    ):
        assert (
            resolve_specialty({"specialty": "", "latest_payload_json": payload})
            == specialty
        )


def test_generic_sheet_marker_is_not_a_specialty():
    assert resolve_specialty({"detail_sheet": "GENERADA"}) == ""
    assert resolve_specialty({"latest_payload_json": "invalid"}) == ""
    assert (
        resolve_specialty({"specialty": "PEDIATRIA", "detail_sheet": "GENERAL"})
        == "PEDIATRIA"
    )


def test_hydration_preserves_specialty():
    from admission_hybrid import AdmissionCloudRepository
    from admission_specialty import with_resolved_specialty

    row = {
        "global_attention_id": "11111111-1111-4111-8111-111111111111",
        "has_detail_sheet": True,
        "latest_payload_json": {"detail_sheet": "PEDIATRIA"},
    }
    event = AdmissionCloudRepository._readthrough_event(row)
    assert event["payload_json"]["detail_sheet"] == "PEDIATRIA"
    assert event["payload_json"]["specialty"] == "PEDIATRIA"
    assert with_resolved_specialty(row)["hoja"] == "PEDIATRIA"
    assert "hoja" not in with_resolved_specialty({"specialty": "PEDIATRIA"})


def test_all_history_cache_preserves_uuid_of_synchronized_attention(tmp_path):
    path = tmp_path / "replica.sqlite"
    global_id = "11111111-1111-4111-8111-111111111111"
    with sqlite3.connect(path) as con:
        con.execute("CREATE TABLE atenciones(id INTEGER,global_attention_id TEXT)")
        con.execute("INSERT INTO atenciones VALUES(458,?)", (global_id,))
    database = SimpleNamespace(
        _connect=lambda: sqlite3.connect(path),
        listar_atenciones=lambda **kwargs: [{"id": 458, "hoja": "PEDIATRIA"}],
    )
    proxy = _HybridDatabaseProxy(database, SimpleNamespace(operational_session=None))
    rows = proxy.list_history_cache_local("listar_atenciones")
    assert rows[0]["global_attention_id"] == global_id


def test_report_totals_and_excel_recover_three_specialties(tmp_path):
    from admission_statistical_reports import build_admission_report_dataset
    from admission_v15_adapter import load_v15_application_module
    from tests.test_admission_statistical_reports_v106 import _filters, _row
    from openpyxl import load_workbook

    rows = [
        _row(index, specialty="") | {"latest_payload_json": {"detail_sheet": specialty}}
        for index, specialty in enumerate(["GENERAL", "PEDIATRIA", "GINECOLOGIA"], 1)
    ]
    dataset = build_admission_report_dataset(rows, _filters())
    counts = _HybridDatabaseProxy._calculate_turn_counts(rows)
    assert counts["GENERAL"] == counts["PEDIATRIA"] == counts["GINECOLOGIA"] == 1
    v15 = load_v15_application_module()
    path = v15.crear_excel_reporte_estadistico(
        dataset.summary, destino=str(tmp_path / "report.xlsx")
    )
    wb = load_workbook(path, read_only=True)
    try:
        assert [wb.active.cell(row, 3).value for row in range(6, 9)] == [
            "GENERAL",
            "PEDIATRIA",
            "GINECOLOGIA",
        ]
    finally:
        wb.close()


def test_bulk_projection_preserves_v15_detail_specialty():
    from admission_hybrid import AdmissionCloudRepository, SyncEvent

    event = SyncEvent(
        "event",
        "attention",
        "11111111-1111-4111-8111-111111111111",
        "CREATE",
        {
            "attention_id": 1,
            "patient_id": 1,
            "detail_sheet": "PEDIATRIA",
            "service_date": "2026-09-06",
        },
        "session",
        1,
        "PC",
        "2026-09-06T10:00:00",
    )
    assert AdmissionCloudRepository._bulk_projection_row(event)[12] == "PEDIATRIA"
