"""Confirm edited exports without changing the registered batch contents."""

import hashlib
import json
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QMessageBox, QWidget

import CALCULOS_QT as app


@pytest.fixture
def exported(tmp_path, monkeypatch):
    path = tmp_path / "listado.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "DATOS ORIGINALES"
    workbook.save(path)
    workbook.close()
    batch = {
        "id": 7,
        "status": app.BATCH_PENDING,
        "period_year": 2026,
        "period_month": 9,
        "ars": "FUTURO",
        "invoice_date": "2026-09-30",
        "invoice_number": "FAC-007",
        "ncf": "B0100000007",
        "provider_code": "14427",
        "provider_name": "HOSPITAL DE PRUEBA",
        "provider_rnc": "430130516",
        "ars_rnc": "101000001",
        "ars_address": "DIRECCION DE PRUEBA",
        "director_name": "DIRECTORA DE PRUEBA",
        "last_export_path": str(path),
    }
    rows = [
        {
            "recibo_id": 13,
            "patient_snapshot": "PACIENTE SINTETICO",
            "document_number_snapshot": "001234567",
            "service_date_snapshot": "2026-09-05",
            "authorization_snapshot": "00001234",
            "total_snapshot": 100.25,
            "estado_facturacion": "FACTURADO",
        }
    ]
    original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    batch["last_export_signature"] = (
        app._monthly_batch_export_signature(batch, rows) + ":" + original_hash
    )
    connection = MagicMock()
    connection.__enter__.return_value = connection
    connection.execute.side_effect = [
        SimpleNamespace(fetchone=lambda: batch),
        SimpleNamespace(fetchone=lambda: None),
        SimpleNamespace(rowcount=1),
        None,
    ]
    monkeypatch.setattr(app, "db_connect", lambda: connection)
    monkeypatch.setattr(app, "_query_monthly_batch_receipts", lambda *_: rows)
    monkeypatch.setattr(app, "get_user", lambda username: USER)
    history = MagicMock()
    monkeypatch.setattr(app, "_insert_action_history", history)
    return SimpleNamespace(
        path=path,
        batch=batch,
        rows=rows,
        original_hash=original_hash,
        connection=connection,
        history=history,
    )


USER = {"username": "tester", "role": app.ROLE_ADMIN}


@pytest.mark.parametrize("change", ["data", "format", "formula"])
def test_edited_excel_can_be_confirmed_and_both_versions_are_audited(exported, change):
    workbook = load_workbook(exported.path)
    if change == "data":
        workbook.active["A1"] = "DATOS CORREGIDOS"
    elif change == "format":
        workbook.active.column_dimensions["A"].width = 40
    else:
        workbook.active["B1"] = "=1+1"
    workbook.save(exported.path)
    workbook.close()

    result = app.confirm_monthly_batch_sent(7, USER)

    assert result["receipts"] == 1
    assert result["total"] == 100.25
    event = json.loads(exported.connection.execute.call_args_list[-1].args[1][-1])
    assert event["excel_modified"] is True
    assert event["exported_file_sha256"] == exported.original_hash
    assert (
        event["sent_file_sha256"]
        == hashlib.sha256(exported.path.read_bytes()).hexdigest()
    )
    assert event["path"] == str(exported.path)
    exported.history.assert_called_once()
    assert exported.rows[0]["patient_snapshot"] == "PACIENTE SINTETICO"


def test_unchanged_excel_still_confirms(exported):
    app.confirm_monthly_batch_sent(7, USER)
    event = json.loads(exported.connection.execute.call_args_list[-1].args[1][-1])
    assert event["excel_modified"] is False
    assert event["sent_file_sha256"] == event["exported_file_sha256"]


@pytest.mark.parametrize("change", ["fiscal", "patient", "total", "membership"])
def test_changes_inside_batch_still_require_a_new_export(exported, change):
    if change == "fiscal":
        exported.batch["invoice_number"] = "FAC-008"
    elif change == "patient":
        exported.rows[0]["patient_snapshot"] = "OTRO PACIENTE"
    elif change == "total":
        exported.rows[0]["total_snapshot"] = 999
    else:
        exported.rows.append({**exported.rows[0], "recibo_id": 14})
    with pytest.raises(ValueError, match="cambió después"):
        app.confirm_monthly_batch_sent(7, USER)
    assert exported.connection.execute.call_count == 1
    exported.history.assert_not_called()


@pytest.mark.parametrize(
    "signature",
    [None, "", "a" * 64, "a:b", "x" * 64 + ":" + "b" * 64, "é:" + "b" * 64],
)
def test_missing_or_malformed_export_proof_is_rejected(exported, signature):
    exported.batch["last_export_signature"] = signature
    with pytest.raises(ValueError, match="cambió después"):
        app.confirm_monthly_batch_sent(7, USER)
    exported.history.assert_not_called()


@pytest.mark.parametrize("path", [None, "", "missing.xlsx"])
def test_missing_export_is_not_confirmed(exported, path):
    exported.batch["last_export_path"] = path
    with pytest.raises(ValueError, match="Excel final"):
        app.confirm_monthly_batch_sent(7, USER)
    exported.history.assert_not_called()


def test_unreadable_excel_does_not_change_state(exported, monkeypatch):
    read = MagicMock(side_effect=PermissionError("Archivo no disponible"))
    monkeypatch.setattr(app, "_file_sha256", read)
    with pytest.raises(PermissionError, match="Archivo no disponible"):
        app.confirm_monthly_batch_sent(7, USER)
    assert exported.connection.execute.call_count == 1
    exported.history.assert_not_called()


def test_permission_is_checked_before_reading_or_writing(exported, monkeypatch):
    monkeypatch.setattr(
        app,
        "get_user",
        lambda _: {
            **USER,
            "permission_overrides": {app.PERMISSION_MANAGE_BILLING_LISTS: False},
        },
    )
    with pytest.raises(PermissionError):
        app.confirm_monthly_batch_sent(7, USER)
    exported.connection.execute.assert_not_called()


@pytest.fixture
def confirmation_view(exported, monkeypatch):
    application = QApplication.instance() or QApplication([])
    view = QWidget()
    view.current_batch_id = 7
    view.current_user = USER
    view.load_batches = MagicMock()
    view.section_tabs = MagicMock()
    monkeypatch.setattr(app, "get_monthly_billing_batch", lambda _: exported.batch)
    monkeypatch.setattr(app, "list_monthly_batch_receipts", lambda _: exported.rows)
    monkeypatch.setattr(app, "FloatingToast", MagicMock())
    yield view
    view.close()
    application.processEvents()


def test_confirmation_explains_which_records_will_be_marked_sent(
    exported, confirmation_view, monkeypatch
):
    question = MagicMock(return_value=QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, "warning", question)
    app.MonthlyBillingListsPage.confirm_current_batch_sent(confirmation_view)
    message = question.call_args.args[2]
    assert "Excel" in message
    assert "no se importan" in message
    assert "Registros: 1" in message
    assert "100.25" in message
    confirmation_view.load_batches.assert_called_once_with(7)
    confirmation_view.section_tabs.setCurrentIndex.assert_called_once_with(1)


def test_cancelling_confirmation_does_not_send(
    exported, confirmation_view, monkeypatch
):
    monkeypatch.setattr(QMessageBox, "warning", lambda *_: QMessageBox.No)
    app.MonthlyBillingListsPage.confirm_current_batch_sent(confirmation_view)
    exported.connection.execute.assert_not_called()
    confirmation_view.load_batches.assert_not_called()
