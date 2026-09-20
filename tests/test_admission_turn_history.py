from datetime import date, datetime
from unittest.mock import Mock
from types import SimpleNamespace

import pytest

from admission_turn_history import (
    historical_turn_period,
    search_turns,
    selected_turn_records,
    turn_cursor,
)
from tests.test_billing_consistency_postgres import SOURCE

pytest_plugins = ["tests.test_billing_consistency_postgres"]


def test_history_adapter_preserves_original_user_and_allows_local_exports(tmp_path):
    from admission_v15_adapter import _HybridDatabaseProxy
    from tests.test_admission_statistical_reports_v106 import (
        _CapturingConnection,
        _LocalDatabase,
        _row,
        _filters,
    )
    from admission_statistical_reports import build_admission_report_dataset
    from ADMISION_PYSIDE6_V15 import facturacion_tabs_pyside6 as v15
    from openpyxl import load_workbook

    record = _row(1)
    record["admission_username"] = "original"
    connection = _CapturingConnection([record])
    runtime = SimpleNamespace(
        offline=False, host=SimpleNamespace(connection_factory=lambda: connection)
    )
    proxy = _HybridDatabaseProxy(_LocalDatabase(), runtime)
    turn = dict(
        operational_source_id=record["operational_source_id"],
        turn_id=316,
        username="original",
        display_name="Admisor original",
    )
    source = proxy.load_historical_turn_report(turn)
    assert source["selected_turn"]["representatives"][0]["username"] == "original"
    assert connection.params[-1] == "original"
    dataset = build_admission_report_dataset(
        source["records"], _filters(), turns=source["turns"]
    )
    pdf, excel = tmp_path / "turn.pdf", tmp_path / "turn.xlsx"
    v15.crear_pdf_reporte(dataset.summary, destino=str(pdf))
    v15.crear_excel_reporte_estadistico(dataset.summary, destino=str(excel))
    assert pdf.read_bytes().startswith(b"%PDF-")
    workbook = load_workbook(excel)
    assert workbook["LISTADO DE PACIENTES"].max_row == 6
    assert workbook["RESUMEN ESTADÍSTICO"]["B2"].value == 1
    workbook.close()
    runtime.offline = True
    with pytest.raises(RuntimeError, match="Conecte"):
        proxy.search_admission_turns(date(2026, 9, 18), date(2026, 9, 18))
    runtime.offline = False
    assert len(proxy.search_admission_turns(date(2026, 9, 18), date(2026, 9, 18))) == 1


def test_turn_history_queries_original_admitter_and_paginates(database):
    with database() as con:
        con.execute(
            "ALTER TABLE admission_operational_sessions ADD COLUMN turn_ends_at TIMESTAMPTZ"
        )
        con.execute(
            "ALTER TABLE admission_operational_turn_intervals ADD COLUMN nominal_ends_at TIMESTAMPTZ, ADD COLUMN active_username TEXT"
        )
        con.execute("CREATE TABLE users(username TEXT UNIQUE,full_name TEXT)")
        con.execute(
            "INSERT INTO users VALUES('original','Admisor original'),('editor','Editor')"
        )
        con.execute(
            "UPDATE admission_operational_sessions SET turn_id=1 WHERE operational_source_id=%s",
            (SOURCE,),
        )
        con.execute("""INSERT INTO admission_operational_turn_intervals(operational_session_id,turn_id,started_at,ended_at,active_username)
            SELECT 'test',n,'2026-09-18 08:00:00-04'::timestamptz + n*interval '1 second',
                '2026-09-19 08:00:00-04','original' FROM generate_series(1,55) n""")
        con.execute(
            "INSERT INTO admission_attention_projection(attention_id,operational_source_id,turn_id,admission_username) VALUES(1,%s,55,'original'),(2,%s,55,'editor')",
            (SOURCE, SOURCE),
        )
        first = search_turns(con, date(2026, 9, 18), date(2026, 9, 18), "original")
        assert len(first) == 51
        assert first[0]["patient_count"] == 1
        second = search_turns(
            con,
            date(2026, 9, 18),
            date(2026, 9, 18),
            "original",
            turn_cursor(first[49]),
        )
        assert len(second) == 5, (
            turn_cursor(first[49]),
            [r["turn_id"] for r in second],
        )
        assert {row["turn_id"] for row in first[:50]}.isdisjoint(
            row["turn_id"] for row in second
        )
        records = selected_turn_records(con, first[0])
        assert [row["attention_id"] for row in records] == [1]
        assert records[0]["admission_username"] == "original"
        assert search_turns(con, date(2026, 9, 19), date(2026, 9, 19)) == []
        assert (
            search_turns(con, date(2026, 9, 18), date(2026, 9, 18), "' OR 1=1 --") == []
        )


def test_invalid_date_range_does_not_query():
    connection = Mock()
    with pytest.raises(ValueError):
        search_turns(connection, date(2026, 9, 19), date(2026, 9, 18))
    connection.execute.assert_not_called()


@pytest.mark.parametrize(
    "start,end",
    [
        (None, None),
        ("bad", "bad"),
        ("2026-09-18", "2026-09-18"),
        ("2026-09-19", "2026-09-18"),
    ],
)
def test_invalid_turn_interval(start, end):
    with pytest.raises(ValueError):
        historical_turn_period({"started_at": start, "ends_at": end})


def test_history_preserves_actual_shift_times():
    result = historical_turn_period(
        {
            "turn_id": 5,
            "display_name": "Original",
            "started_at": "2026-09-18T12:35:00+00:00",
            "ends_at": "2026-09-19T13:00:00+00:00",
        }
    )
    assert result.start_at == datetime(2026, 9, 18, 8, 35)
    assert result.end_at == datetime(2026, 9, 19, 9)
