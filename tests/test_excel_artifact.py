"""Filesystem integration tests with synthetic workbooks only."""

from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from pathlib import Path
from unittest.mock import Mock
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

import excel_artifact as artifacts


def write_value(target, value):
    book = Workbook()
    book.active["A1"] = value
    book.active["B2"] = "=1+1"
    book.active.merge_cells("C3:D3")
    try:
        artifacts.save_workbook(book, target)
    finally:
        book.close()


def read_value(path):
    book = load_workbook(path)
    try:
        assert book.active["B2"].value == "=1+1"
        assert "C3:D3" in book.active.merged_cells
        return book.active["A1"].value
    finally:
        book.close()


@pytest.mark.parametrize("executor", [ThreadPoolExecutor, ProcessPoolExecutor])
def test_concurrent_writers_keep_complete_artifacts_and_previous_version(
    tmp_path, executor
):
    target = tmp_path / "concurrent.xlsx"
    with executor(max_workers=3) as pool:
        list(pool.map(write_value, [target] * 6, range(6)))
    assert read_value(target) in range(6)
    assert read_value(tmp_path / "concurrent.last-valid.xlsx") in range(6)
    assert not list(tmp_path.glob("*.pending.xlsx"))


def test_copy_preserves_exact_bytes_and_last_valid(tmp_path):
    source, target = tmp_path / "source.xlsx", tmp_path / "target.xlsx"
    write_value(source, "new")
    write_value(target, "previous")
    previous = target.read_bytes()
    artifacts.copy_workbook(source, target)
    assert target.read_bytes() == source.read_bytes()
    assert (tmp_path / "target.last-valid.xlsx").read_bytes() == previous


@pytest.mark.parametrize(
    "failure", ["write", "invalid_zip", "invalid_workbook", "replace", "backup"]
)
def test_failed_publish_preserves_previous_and_cleans_temporary(
    tmp_path, monkeypatch, failure
):
    target = tmp_path / "target.xlsx"
    write_value(target, "previous")
    previous = target.read_bytes()
    book = Workbook()
    if failure == "write":
        book.save = Mock(side_effect=OSError("disk failure"))
    elif failure == "invalid_zip":
        book.save = lambda path: Path(path).write_bytes(b"incomplete zip")
    elif failure == "invalid_workbook":

        def invalid_workbook(path):
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("invalid.xml", "<invalid/>")

        book.save = invalid_workbook
    else:
        real_replace = artifacts.os.replace

        def fail_replace(source, destination):
            if failure == "backup" or Path(destination) == target:
                raise PermissionError("file in use")
            real_replace(source, destination)

        monkeypatch.setattr(artifacts.os, "replace", fail_replace)
    with pytest.raises((OSError, zipfile.BadZipFile, KeyError)):
        artifacts.save_workbook(book, target)
    book.close()
    assert target.read_bytes() == previous
    assert not list(tmp_path.glob("*.pending.xlsx"))


def test_corrupt_destination_is_not_overwritten_and_backup_is_preserved(tmp_path):
    target = tmp_path / "target.xlsx"
    write_value(target, "first")
    write_value(target, "second")
    backup = (tmp_path / "target.last-valid.xlsx").read_bytes()
    target.write_bytes(b"evidence")
    with pytest.raises(zipfile.BadZipFile):
        write_value(target, "third")
    assert target.read_bytes() == b"evidence"
    assert (tmp_path / "target.last-valid.xlsx").read_bytes() == backup


def test_crc_failure_is_rejected(tmp_path, monkeypatch):
    target = tmp_path / "target.xlsx"
    write_value(target, "valid")
    monkeypatch.setattr(zipfile.ZipFile, "testzip", lambda _self: "bad.xml")
    with pytest.raises(zipfile.BadZipFile, match="CRC"):
        artifacts.validate_xlsx(target)


def test_lock_release_after_exception_allows_retry(tmp_path):
    target = tmp_path / "target.xlsx"
    with pytest.raises(RuntimeError), artifacts.artifact_lock(target):
        raise RuntimeError("failure")
    write_value(target, "retry")
    assert read_value(target) == "retry"
