import json
from datetime import datetime, timezone
from uuid import uuid4

from openpyxl import load_workbook

import CALCULOS_QT as app
import report_documents
from billing_closure_recovery import closure_from_interval
from tests import test_integral_emergency_to_monthly_list as integration


def test_central_capture_header_pdf_data_and_excel_agree(tmp_path):
    fixture = integration.IntegralEmergencyToMonthlyListTests(
        "test_emergency_to_audit_to_monthly_ars_list"
    )
    fixture.setUp()
    try:
        source, session, transition = (str(uuid4()) for _ in range(3))
        start = datetime(2026, 9, 24, 12, tzinfo=timezone.utc)
        end = datetime(2026, 9, 25, 12, tzinfo=timezone.utc)
        with app.db_connect() as con:
            con.execute(
                """INSERT INTO admission_operational_sessions(
                operational_session_id,operational_source_id,active_username,
                primary_device_id,primary_login_session_id,turn_id)
                VALUES(%s,%s,'SYNTHETIC','TEST','TEST',11)""",
                (session, source),
            )
            con.execute(
                """INSERT INTO admission_operational_turn_intervals(
                operational_session_id,generation,turn_id,active_username,started_at,ended_at)
                VALUES(%s,1,10,'SYNTHETIC',%s,%s)""",
                (session, start, end),
            )
            con.execute("""INSERT INTO recibos(numero,created_at,total,numero_autorizacion,ars)
                SELECT n,'2026-09-24T15:00:00-04',6.00,'1234','HUMANO'
                FROM generate_series(1,50) n""")
            details = {
                "status": "COMMITTED",
                "request": {
                    "operational_source_id": source,
                    "transition_type": "PRIMARY_USER_HANDOFF",
                },
                "result": {"old_turn_id": 10, "new_turn_id": 11},
            }
            con.execute(
                """INSERT INTO admission_operational_audit(
                operational_session_id,event_type,username,details_json,transition_id)
                VALUES(%s,'TURN_HANDOFF_TRANSITION','SYNTHETIC',%s::jsonb,%s)""",
                (session, json.dumps(details), transition),
            )
        event = closure_from_interval(
            {
                "operational_source_id": source,
                "operational_session_id": session,
                "turn_id": 10,
                "started_at": start,
                "ended_at": end,
                "active_username": "SYNTHETIC",
            }
        )
        header = app.capture_shift_closure_snapshot(event, [])
        assert header["classification_version"] == 3
        assert header["worked_applicable"] == 50
        assert not app.capture_shift_closure_snapshot(event, [])["snapshot_created"]
        data = app.build_shift_closure_report_data(header)
        assert data["receipt_count"] == data["historical_billed"] == 50
        assert data["amount"] == 300
        app.save_shift_report_snapshot(
            header,
            data,
            {"mode": "shift_closure", "data": data, "generated_by": "SYNTHETIC"},
            end.isoformat(),
        )
        with app.db_connect() as con:
            record = report_documents.load_current_report_snapshot(
                con, "billing_shift_closures", f"{source}|10"
            )
        path = tmp_path / "close.xlsx"
        report_documents.export_report_snapshot_xlsx(record, str(path))
        workbook = load_workbook(path, data_only=True)
        values = dict(
            workbook["Resumen"].iter_rows(min_row=11, max_col=2, values_only=True)
        )
        assert values["Recibos guardados durante el turno"] == 50
        assert values["Importe de recibos del turno"] == 300
        assert "receipt_ids" not in str(list(workbook["Datos históricos"].values))
        workbook.close()
    finally:
        fixture.tearDown()
